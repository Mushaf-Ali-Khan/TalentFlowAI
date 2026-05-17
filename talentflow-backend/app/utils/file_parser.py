import io
import logging
import pdfplumber
import fitz  # PyMuPDF
import docx
from typing import Tuple, Optional
# Note: pytesseract and pdf2image would be required for OCR but are optional and require system deps.
# Using a placeholder for OCR as defined in the spec.

logger = logging.getLogger(__name__)

def parse_pdf_pdfplumber(file_bytes: bytes) -> Tuple[Optional[str], float]:
    """Strategy 1: Use pdfplumber"""
    try:
        text = ""
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
        
        confidence = 1.0 if len(text) > 200 else 0.5
        if not text.strip():
            return None, 0.0
        return text, confidence
    except Exception as e:
        logger.error(f"pdfplumber failed: {e}")
        return None, 0.0

def parse_pdf_pymupdf(file_bytes: bytes) -> Tuple[Optional[str], float]:
    """Strategy 2: Use PyMuPDF as fallback"""
    try:
        text = ""
        with fitz.open(stream=file_bytes, filetype="pdf") as doc:
            for page in doc:
                text += page.get_text() + "\n"
                
        confidence = 0.8 if len(text) > 200 else 0.4
        if not text.strip():
            return None, 0.0
        return text, confidence
    except Exception as e:
        logger.error(f"pymupdf failed: {e}")
        return None, 0.0

def parse_pdf_ocr(file_bytes: bytes) -> Tuple[Optional[str], float]:
    """Strategy 3: OCR fallback"""
    # Placeholder for actual OCR implementation which requires tesseract installed on OS
    # For now, it returns None to simulate failure or lack of OCR support
    logger.warning("OCR fallback invoked but not fully implemented.")
    return None, 0.0

def parse_docx(file_bytes: bytes) -> Tuple[Optional[str], float]:
    """Strategy for DOCX files"""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        text = "\n".join([para.text for para in doc.paragraphs])
        
        confidence = 1.0 if len(doc.paragraphs) > 0 else 0.0
        if not text.strip():
            return None, 0.0
        return text, confidence
    except Exception as e:
        logger.error(f"docx parsing failed: {e}")
        return None, 0.0

def extract_text_from_file(file_bytes: bytes, file_type: str) -> Tuple[Optional[str], float, Optional[str]]:
    """Returns (text, confidence, strategy_used)"""
    file_type = file_type.lower().lstrip('.')
    
    if file_type == 'docx' or file_type == 'doc':
        text, conf = parse_docx(file_bytes)
        if text:
            return text, conf, 'docx'
        return None, 0.0, None
        
    elif file_type == 'pdf':
        text, conf = parse_pdf_pdfplumber(file_bytes)
        if text and conf > 0.5:
            return text, conf, 'pdfplumber'
            
        logger.info("Falling back to pymupdf")
        text, conf = parse_pdf_pymupdf(file_bytes)
        if text and conf > 0.4:
            return text, conf, 'pymupdf'
            
        logger.info("Falling back to OCR")
        text, conf = parse_pdf_ocr(file_bytes)
        if text:
            return text, conf, 'ocr'
            
    return None, 0.0, None
